"""社区提示词 API"""
from __future__ import annotations

import csv
import io

from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Query
from sqlalchemy import select, func, or_
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.deps import get_current_user
from app.db.session import get_db
from app.adapters.storage import storage
from app.models.prompt import PromptCategory, PromptExample
from app.models.prompt_moderation import PromptReport, PromptAuditLog
from app.models.user import User
from app.schemas.common import ApiResponse

router = APIRouter()

MAX_UPLOAD_SIZE = 10 * 1024 * 1024
ALLOWED_IMAGE_TYPES = {"image/jpeg", "image/png", "image/webp", "image/gif", "image/svg+xml"}
REPORT_REASONS = {
    "违规内容",
    "侵权内容",
    "垃圾广告",
    "不实信息",
    "低质重复",
    "其他",
}


def _is_admin(user: User) -> bool:
    return user.role == "admin"


def _normalize_report_reason(raw: str | None) -> str:
    reason = (raw or "其他").strip() or "其他"
    if reason not in REPORT_REASONS:
        raise HTTPException(status_code=400, detail=f"不支持的举报原因: {reason}")
    return reason


async def _get_or_create_category(db: AsyncSession, name: str) -> PromptCategory:
    category_name = (name or "").strip() or "未分类"
    result = await db.execute(
        select(PromptCategory).where(
            PromptCategory.name == category_name,
            PromptCategory.deleted_at.is_(None),
        )
    )
    category = result.scalar_one_or_none()
    if category:
        return category

    category = PromptCategory(name=category_name, sort_order=0)
    db.add(category)
    await db.flush()
    return category


@router.get("")
async def get_prompts(
    keyword: str | None = None,
    category: str | None = None,
    owner: bool = Query(False, description="仅查看自己上传的提示词"),
    page: int = Query(1, ge=1),
    limit: int = Query(100, ge=1, le=1000),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """获取社区提示词列表"""
    conditions = [
        PromptExample.deleted_at.is_(None),
        PromptCategory.deleted_at.is_(None),
    ]

    if owner:
        conditions.append(PromptExample.created_by == current_user.id)
    elif not _is_admin(current_user):
        # 社区页默认仅展示公开内容（管理员可见全部）
        conditions.append(PromptExample.is_public.is_(True))

    if category:
        conditions.append(PromptCategory.name == category)

    if keyword:
        kw = f"%{keyword}%"
        conditions.append(
            or_(
                PromptExample.name.ilike(kw),
                PromptExample.chinese_example.ilike(kw),
                PromptExample.english_example.ilike(kw),
                PromptExample.param_type.ilike(kw),
                PromptCategory.name.ilike(kw),
            )
        )

    base_stmt = (
        select(PromptExample, PromptCategory)
        .join(PromptCategory, PromptCategory.id == PromptExample.category_id)
        .where(*conditions)
    )

    count_stmt = select(func.count()).select_from(
        base_stmt.order_by(None).subquery()
    )
    total_result = await db.execute(count_stmt)
    total = total_result.scalar() or 0

    result = await db.execute(
        base_stmt
        .order_by(PromptExample.created_at.desc(), PromptExample.id.desc())
        .offset((page - 1) * limit)
        .limit(limit)
    )
    rows = result.all()

    creator_ids = {example.created_by for example, _ in rows if example.created_by is not None}
    creator_map: dict[int, str] = {}
    if creator_ids:
        users_result = await db.execute(
            select(User.id, User.username).where(User.id.in_(creator_ids))
        )
        creator_map = {uid: uname for uid, uname in users_result.all()}

    is_admin = _is_admin(current_user)
    items = []
    for example, cat in rows:
        can_edit = is_admin or (example.created_by == current_user.id)
        title = (example.name or "").strip()
        if not title:
            title = f"{cat.name} - {example.param_type or '通用'}"

        items.append(
            {
                "id": example.id,
                "title": title,
                "name": example.name,
                "chinese": example.chinese_example,
                "english": example.english_example,
                "image_url": example.image_url or "",
                "category": cat.name,
                "param_type": example.param_type or "",
                "created_by": example.created_by,
                "created_by_name": creator_map.get(example.created_by, "系统") if example.created_by else "系统",
                "is_public": example.is_public,
                "can_edit": can_edit,
                "can_delete": can_edit,
                "is_mine": example.created_by == current_user.id,
                "created_at": str(example.created_at),
            }
        )

    return ApiResponse(
        data={
            "items": items,
            "total": total,
            "page": page,
            "limit": limit,
        }
    )


@router.get("/categories")
async def get_categories(
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """获取提示词分类统计"""
    conditions = [
        PromptCategory.deleted_at.is_(None),
        PromptExample.deleted_at.is_(None),
    ]
    if not _is_admin(current_user):
        conditions.append(PromptExample.is_public.is_(True))

    result = await db.execute(
        select(PromptCategory.name, func.count(PromptExample.id))
        .join(PromptExample, PromptExample.category_id == PromptCategory.id)
        .where(*conditions)
        .group_by(PromptCategory.id, PromptCategory.name)
        .order_by(func.count(PromptExample.id).desc(), PromptCategory.id.asc())
    )
    categories = [{"name": name, "count": count} for name, count in result.all()]
    return ApiResponse(data={"categories": categories})


@router.post("")
async def create_prompt(
    data: dict,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """添加社区提示词（默认归属当前用户）"""
    chinese = (data.get("chinese") or "").strip()
    if not chinese:
        raise HTTPException(status_code=400, detail="中文提示词不能为空")

    category = await _get_or_create_category(db, data.get("category") or "未分类")
    prompt_name = (data.get("name") or data.get("title") or "").strip() or None
    example = PromptExample(
        category_id=category.id,
        param_type=(data.get("param_type") or "").strip() or "通用",
        image_url=(data.get("image_url") or "").strip() or None,
        name=prompt_name,
        chinese_example=chinese,
        english_example=(data.get("english") or "").strip(),
        ul_list=data.get("ul_list", []),
        sort_order=int(data.get("sort_order", 0) or 0),
        is_public=bool(data.get("is_public", True)),
        created_by=current_user.id,
        updated_by=current_user.id,
    )
    db.add(example)
    await db.commit()
    await db.refresh(example)

    return ApiResponse(
        data={
            "id": example.id,
            "title": prompt_name or f"{category.name} - {example.param_type}",
            "name": example.name,
            "chinese": example.chinese_example,
            "english": example.english_example,
            "image_url": example.image_url or "",
            "category": category.name,
            "param_type": example.param_type,
            "created_by": example.created_by,
            "created_by_name": current_user.username,
            "is_public": example.is_public,
            "can_edit": True,
            "can_delete": True,
            "is_mine": True,
        }
    )


@router.put("/{prompt_id}")
async def update_prompt(
    prompt_id: int,
    data: dict,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """编辑提示词（仅本人或管理员）"""
    result = await db.execute(
        select(PromptExample, PromptCategory)
        .join(PromptCategory, PromptCategory.id == PromptExample.category_id)
        .where(
            PromptExample.id == prompt_id,
            PromptExample.deleted_at.is_(None),
            PromptCategory.deleted_at.is_(None),
        )
    )
    row = result.first()
    if not row:
        raise HTTPException(status_code=404, detail="提示词不存在")

    example, current_category = row
    if not (_is_admin(current_user) or example.created_by == current_user.id):
        raise HTTPException(status_code=403, detail="无权编辑该提示词")

    if "category" in data:
        new_category = await _get_or_create_category(db, data.get("category") or "未分类")
        example.category_id = new_category.id
        current_category = new_category

    if "chinese" in data:
        chinese = (data.get("chinese") or "").strip()
        if not chinese:
            raise HTTPException(status_code=400, detail="中文提示词不能为空")
        example.chinese_example = chinese

    if "english" in data:
        example.english_example = (data.get("english") or "").strip()
    if "image_url" in data:
        example.image_url = (data.get("image_url") or "").strip() or None
    if "param_type" in data:
        example.param_type = (data.get("param_type") or "").strip() or "通用"
    if "name" in data or "title" in data:
        example.name = (data.get("name") or data.get("title") or "").strip() or None
    if "is_public" in data and _is_admin(current_user):
        example.is_public = bool(data.get("is_public"))

    example.updated_by = current_user.id
    await db.commit()
    await db.refresh(example)

    return ApiResponse(
        data={
            "id": example.id,
            "title": (example.name or f"{current_category.name} - {example.param_type}"),
            "name": example.name,
            "chinese": example.chinese_example,
            "english": example.english_example,
            "image_url": example.image_url or "",
            "category": current_category.name,
            "param_type": example.param_type,
            "created_by": example.created_by,
            "created_by_name": current_user.username if example.created_by == current_user.id else None,
            "is_public": example.is_public,
            "can_edit": True,
            "can_delete": True,
            "is_mine": example.created_by == current_user.id,
        }
    )


@router.delete("/{prompt_id}")
async def delete_prompt(
    prompt_id: int,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """删除提示词（仅本人或管理员）"""
    result = await db.execute(
        select(PromptExample).where(
            PromptExample.id == prompt_id,
            PromptExample.deleted_at.is_(None),
        )
    )
    example = result.scalar_one_or_none()
    if not example:
        raise HTTPException(status_code=404, detail="提示词不存在")

    if not (_is_admin(current_user) or example.created_by == current_user.id):
        raise HTTPException(status_code=403, detail="无权删除该提示词")

    example.deleted_at = func.now()
    example.updated_by = current_user.id
    await db.commit()
    return ApiResponse(message="已删除")


@router.post("/import")
async def import_prompts(
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """从 Excel/CSV 批量导入提示词"""
    import openpyxl

    content = await file.read()
    filename = (file.filename or "").lower()

    rows: list[dict] = []
    failed_rows: list[dict] = []
    if filename.endswith(".csv"):
        text = content.decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(text))
        for row_no, row in enumerate(reader, start=2):
            title = (row.get("标题") or row.get("title") or "").strip()
            chinese = (row.get("中文提示词") or row.get("中文") or row.get("chinese") or "").strip()
            english = (row.get("英文提示词") or row.get("英文") or row.get("english") or "").strip()
            category = (row.get("分类") or row.get("category") or "").strip() or "未分类"
            image_url = (row.get("图片") or row.get("图片URL") or row.get("image_url") or "").strip()
            param_type = (row.get("参数类型") or row.get("param_type") or "").strip() or "通用"
            if not chinese:
                continue
            if not image_url:
                failed_rows.append({
                    "row": row_no,
                    "title": title,
                    "reason": "图片URL不能为空",
                })
                continue
            rows.append(
                {
                    "title": title,
                    "chinese": chinese,
                    "english": english,
                    "category": category,
                    "image_url": image_url,
                    "param_type": param_type,
                }
            )
    elif filename.endswith((".xlsx", ".xls")):
        wb = openpyxl.load_workbook(filename=io.BytesIO(content))
        ws = wb.active
        headers = [str(cell.value or "").strip() for cell in ws[1]]

        col_map: dict[str, int] = {}
        for i, h in enumerate(headers):
            if h in {"标题", "title", "Title"}:
                col_map["title"] = i
            elif h in {"中文提示词", "中文", "chinese", "Chinese"}:
                col_map["chinese"] = i
            elif h in {"英文提示词", "英文", "english", "English"}:
                col_map["english"] = i
            elif h in {"分类", "category", "Category"}:
                col_map["category"] = i
            elif h in {"参数类型", "param_type", "ParamType"}:
                col_map["param_type"] = i
            elif h in {"图片", "图片URL", "image_url", "ImageURL"}:
                col_map["image_url"] = i

        for row_no, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            chinese = str(row[col_map.get("chinese", 1)] or "").strip()
            if not chinese:
                continue
            image_url = str(row[col_map.get("image_url", 5)] or "").strip()
            title = str(row[col_map.get("title", 0)] or "").strip()
            if not image_url:
                failed_rows.append({
                    "row": row_no,
                    "title": title,
                    "reason": "图片URL不能为空",
                })
                continue
            rows.append(
                {
                    "title": title,
                    "chinese": chinese,
                    "english": str(row[col_map.get("english", 2)] or "").strip(),
                    "category": str(row[col_map.get("category", 3)] or "").strip() or "未分类",
                    "param_type": str(row[col_map.get("param_type", 4)] or "").strip() or "通用",
                    "image_url": image_url,
                }
            )
    else:
        raise HTTPException(status_code=400, detail="仅支持 .xlsx / .xls / .csv 文件")

    total_rows = len(rows) + len(failed_rows)
    if total_rows == 0:
        raise HTTPException(status_code=400, detail="未找到有效数据，请检查文件格式和表头")
    if not rows:
        return ApiResponse(
            message="导入完成（0 成功）",
            data={
                "total_rows": total_rows,
                "imported_count": 0,
                "failed_count": len(failed_rows),
                "failed_rows": failed_rows[:200],
                "created_categories": [],
            },
        )

    categories = [row["category"] for row in rows]
    existing_result = await db.execute(
        select(PromptCategory).where(
            PromptCategory.name.in_(categories),
            PromptCategory.deleted_at.is_(None),
        )
    )
    category_map = {cat.name: cat for cat in existing_result.scalars().all()}

    new_category_names = [name for name in set(categories) if name not in category_map]
    created_categories: list[str] = []
    for name in new_category_names:
        cat = PromptCategory(name=name)
        db.add(cat)
        await db.flush()
        category_map[name] = cat
        created_categories.append(name)

    for row in rows:
        cat = category_map[row["category"]]
        db.add(
            PromptExample(
                category_id=cat.id,
                param_type=row["param_type"],
                image_url=row["image_url"] or None,
                name=row["title"] or None,
                chinese_example=row["chinese"],
                english_example=row["english"],
                ul_list=[],
                created_by=current_user.id,
                updated_by=current_user.id,
            )
        )

    await db.commit()
    imported_count = len(rows)
    failed_count = len(failed_rows)
    msg = f"导入完成：成功 {imported_count} 条，失败 {failed_count} 条"
    return ApiResponse(
        message=msg,
        data={
            "total_rows": total_rows,
            "imported_count": imported_count,
            "failed_count": failed_count,
            "failed_rows": failed_rows[:200],
            "created_categories": created_categories,
        },
    )


@router.post("/upload-image")
async def upload_prompt_image(
    file: UploadFile = File(...),
    current_user: User = Depends(get_current_user),
):
    """上传提示词示例图片，返回可访问 URL"""
    _ = current_user
    content = await file.read()
    if len(content) > MAX_UPLOAD_SIZE:
        raise HTTPException(status_code=400, detail="图片大小超过 10MB 限制")

    content_type = (file.content_type or "").lower()
    if content_type not in ALLOWED_IMAGE_TYPES:
        raise HTTPException(status_code=400, detail=f"不支持的图片类型: {content_type or 'unknown'}")

    url = await storage.save(
        file_content=content,
        filename=file.filename or "prompt-image.png",
        content_type=content_type,
        subdir="prompts",
    )
    return ApiResponse(data={"url": url})


@router.post("/{prompt_id}/report")
async def report_prompt(
    prompt_id: int,
    data: dict,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """举报提示词（社区用户）"""
    result = await db.execute(
        select(PromptExample).where(
            PromptExample.id == prompt_id,
            PromptExample.deleted_at.is_(None),
        )
    )
    prompt = result.scalar_one_or_none()
    if not prompt:
        raise HTTPException(status_code=404, detail="提示词不存在")

    existing = await db.execute(
        select(PromptReport).where(
            PromptReport.prompt_id == prompt_id,
            PromptReport.reporter_id == current_user.id,
            PromptReport.status == "pending",
        )
    )
    if existing.scalar_one_or_none():
        raise HTTPException(status_code=400, detail="你已举报过该提示词，请等待处理")

    reason = _normalize_report_reason(data.get("reason"))
    details = (data.get("details") or "").strip() or None
    report = PromptReport(
        prompt_id=prompt_id,
        reporter_id=current_user.id,
        reason=reason,
        details=details,
        status="pending",
    )
    db.add(report)
    db.add(
        PromptAuditLog(
            prompt_id=prompt_id,
            action="report",
            operator_id=current_user.id,
            details=f"reason={reason}; details={details or ''}",
        )
    )
    await db.commit()
    return ApiResponse(message="举报已提交，管理员将尽快处理")


@router.get("/report-reasons")
async def report_reasons(
    current_user: User = Depends(get_current_user),
):
    _ = current_user
    return ApiResponse(data={"reasons": sorted(REPORT_REASONS)})


@router.get("/my-reports")
async def my_reports(
    page: int = Query(1, ge=1),
    limit: int = Query(20, ge=1, le=100),
    status: str | None = Query(default=None),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """查看当前用户提交的举报记录"""
    conditions = [PromptReport.reporter_id == current_user.id]
    if status:
        conditions.append(PromptReport.status == status)

    count_stmt = select(func.count()).select_from(PromptReport).where(*conditions)
    total_result = await db.execute(count_stmt)
    total = total_result.scalar() or 0

    stmt = (
        select(PromptReport, PromptExample)
        .join(PromptExample, PromptExample.id == PromptReport.prompt_id)
        .where(*conditions)
        .order_by(PromptReport.created_at.desc())
        .offset((page - 1) * limit)
        .limit(limit)
    )
    result = await db.execute(stmt)
    rows = result.all()

    return ApiResponse(data={
        "items": [
            {
                "id": report.id,
                "prompt_id": report.prompt_id,
                "prompt_name": prompt.name or f"#{prompt.id}",
                "prompt_text": prompt.chinese_example[:80],
                "reason": report.reason,
                "details": report.details,
                "status": report.status,
                "created_at": str(report.created_at),
                "resolved_at": str(report.resolved_at) if report.resolved_at else None,
                "resolution_note": report.resolution_note,
            }
            for report, prompt in rows
        ],
        "total": total,
        "page": page,
        "limit": limit,
    })
